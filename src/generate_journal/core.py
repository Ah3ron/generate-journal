"""Генерация журналов посещаемости в формате .docx."""

from __future__ import annotations

import argparse
import calendar
import sys
from pathlib import Path
from typing import TYPE_CHECKING, cast

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Length, Pt

if TYPE_CHECKING:
    from collections.abc import Sequence

    # `docx.Document` — это фабричная функция, а не класс; настоящий класс
    # живёт в `docx.document` и нужен только для аннотаций типов.
    from docx.document import Document as _Document
    from docx.styles.style import _ParagraphStyle

DEFAULT_PERIOD = "октябрь 2026"
DEFAULT_STUDENTS_FILE = "ученики.txt"
DEFAULT_EXCLUDE = ""
DEFAULT_OUTPUT_DIR = Path("output")

WEEKEND = (5, 6)  # 5=сб, 6=вс; если суббота учебная: (6,)
BORDER_THIN = "8"  # обычная линия (1 pt)
BORDER_THICK = "18"  # жирная линия блока ИТОГ (2.25 pt)
MIN_EMPTY_ROWS = 22  # минимум строк под список учеников

MONTHS = {
    "январь": 1,
    "февраль": 2,
    "март": 3,
    "апрель": 4,
    "май": 5,
    "июнь": 6,
    "июль": 7,
    "август": 8,
    "сентябрь": 9,
    "октябрь": 10,
    "ноябрь": 11,
    "декабрь": 12,
}

ITOG_LABELS = (
    "болен - б",
    "по заявлению родителей - з",
    "санаторий - с",
    "по неуважительной причине - н",
    "Всего отсутствующих -",
)

ClassInfo = tuple[str, str, list[str]]  # (метка класса, учитель, имена)
Journal = tuple[str, list[ClassInfo]]  # (название журнала, классы)


def make_dates(
    period: str,
    exclude: frozenset[str] = frozenset(),
    weekend: tuple[int, ...] = WEEKEND,
) -> list[str]:
    """Учебные дни месяца: 'октябрь 2026' -> ['01.10', '02.10', ...]."""
    try:
        name, year = period.strip().rsplit(" ", 1)
        month = MONTHS[name.lower()]
    except ValueError, KeyError:
        raise SystemExit(
            f"Не удалось разобрать период: {period!r} (пример: 'октябрь 2026')"
        ) from None
    out: list[str] = []
    for d in calendar.Calendar(firstweekday=0).itermonthdates(int(year), month):
        if d.month != month or d.weekday() in weekend:
            continue
        day = f"{d.day:02d}.{d.month:02d}"
        if day not in exclude:
            out.append(day)
    return out


def load_students(path: str) -> list[Journal]:
    """Читает файл: @ЖУРНАЛ — название, @КЛАСС метка | учитель, дальше имена."""
    journals: list[Journal] = []
    journal: Journal | None = None
    cls: ClassInfo | None = None
    with open(path, encoding="utf-8-sig") as f:
        for raw in f:
            line = raw.strip().rstrip(",")
            if not line or line.startswith("#"):
                continue
            up = line.upper()
            if up.startswith("@ЖУРНАЛ"):
                journal = (line[7:].strip(), [])
                journals.append(journal)
                cls = None
            elif up.startswith("@КЛАСС"):
                if journal is None:
                    journal = ("Без названия", [])
                    journals.append(journal)
                label, _, teacher = line[6:].strip().partition("|")
                cls = (label.strip(), teacher.strip(), [])
                journal[1].append(cls)
            elif cls is not None:
                cls[2].extend(s.strip() for s in line.split(",") if s.strip())
    return journals


def _set_borders(cell, sz: str = BORDER_THIN) -> None:
    """Явные рамки ячейки (старые определения заменяются)."""
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("w:tcBorders")):
        tcPr.remove(old)
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        e = OxmlElement("w:" + edge)
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), sz)
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), "000000")
        borders.append(e)
    tcPr.append(borders)


def _rotate(cell) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    td = OxmlElement("w:textDirection")
    td.set(qn("w:val"), "btLr")
    tcPr.append(td)


def _tight(cell, center: bool = False, size: int | None = None) -> None:
    for p in cell.paragraphs:
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        if center:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if size:
            for r in p.runs:
                r.font.size = Pt(size)


def _add_class(
    doc: _Document,
    label: str,
    teacher: str,
    names: list[str],
    dates: list[str],
    widths: Sequence[Length],
    first: bool = False,
) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    if not first:
        p.add_run().add_break(WD_BREAK.PAGE)
    r = p.add_run(f"{label} класс, {teacher}")
    r.bold = True
    r.font.size = Pt(12)

    t = doc.add_table(rows=1, cols=len(dates) + 3)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    t._tbl.tblPr.append(lay)

    hdr = t.rows[0]
    hdr.height = Cm(1.8)
    hdr.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
    c = hdr.cells[0]
    c.paragraphs[0].add_run("№ п/п").bold = True
    _tight(c, center=True, size=9)
    c = hdr.cells[1]
    pr = c.paragraphs[0]
    r1 = pr.add_run("ФИО")
    r1.bold = True
    r1.add_break()
    r2 = pr.add_run(label)
    r2.bold = True
    _tight(c, center=True)
    for i, d in enumerate(dates):
        c = hdr.cells[2 + i]
        run = c.paragraphs[0].add_run(d)
        run.bold = True
        run.font.size = Pt(9)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        _rotate(c)

    for i in range(max(len(names), MIN_EMPTY_ROWS)):
        row = t.add_row()
        row.height = Cm(0.5)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        row.cells[0].text = f"{i + 1}." if i < len(names) else ""
        row.cells[1].text = names[i] if i < len(names) else ""
        _tight(row.cells[0], center=True)
        _tight(row.cells[1])
        for c in row.cells:
            c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    fi = len(t.rows)  # начало блока ИТОГ
    for lab in ITOG_LABELS:
        row = t.add_row()
        row.height = Cm(0.45)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        run = row.cells[1].paragraphs[0].add_run(lab)
        run.bold = True
        run.font.size = Pt(9)
        _tight(row.cells[1])
        for c in row.cells:
            c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    mc = t.cell(fi, 0).merge(t.cell(fi + 4, 0))
    mp = mc.paragraphs[0]
    mp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for k, ch in enumerate("ИТОГ"):
        r = mp.add_run(ch)
        r.bold = True
        r.font.size = Pt(9)
        if k < 3:
            r.add_break()

    for r_idx, row in enumerate(t.rows):  # ИТОГ — жирная рамка, остальное — тонкая
        sz = BORDER_THICK if r_idx >= fi else BORDER_THIN
        for i, c in enumerate(row.cells):
            _set_borders(c, sz)
            if i < len(widths):
                c.width = widths[i]


def make_journal(
    jname: str,
    classes: list[ClassInfo],
    period: str,
    dates: list[str],
    output_dir: Path,
) -> None:
    fname = output_dir / f"Журнал посещаемости {jname}, {period}.docx"
    doc = Document()
    normal = cast("_ParagraphStyle", doc.styles["Normal"])
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10)
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(1.0)
    sec.top_margin = sec.bottom_margin = Cm(0.9)
    widths = [Cm(0.9), Cm(8.4)] + [Cm(0.8)] * len(dates) + [Cm(0.6)]
    for k, (label, teacher, names) in enumerate(classes):
        _add_class(doc, label, teacher, names, dates, widths, first=(k == 0))
        print(f"   {label} — {len(names)} чел.")
    doc.save(str(fname))
    print(f"Создан: {fname}")


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Генерация журналов посещаемости (docx)"
    )
    parser.add_argument(
        "-p",
        "--period",
        default=DEFAULT_PERIOD,
        help="месяц и год, например: 'октябрь 2026' (по умолчанию: %(default)s)",
    )
    parser.add_argument(
        "-s",
        "--students",
        default=DEFAULT_STUDENTS_FILE,
        help="файл со списками классов (по умолчанию: %(default)s)",
    )
    parser.add_argument(
        "-e",
        "--exclude",
        default=DEFAULT_EXCLUDE,
        help="нерабочие дни через запятую: 05.10,06.10",
    )
    parser.add_argument(
        "-o",
        "--output-dir",
        type=Path,
        default=DEFAULT_OUTPUT_DIR,
        help="каталог для сохранения .docx-журналов (по умолчанию: %(default)s)",
    )
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    exclude = frozenset(s.strip() for s in args.exclude.split(",") if s.strip())
    dates = make_dates(args.period, exclude)
    print(f"Учебные дни {args.period} ({len(dates)}): {' '.join(dates)}")

    students_path = Path(args.students)
    if not students_path.exists():
        print(
            f"Не найден файл '{students_path}' — положите его рядом и запустите снова.",
            file=sys.stderr,
        )
        return 1

    output_dir: Path = args.output_dir
    output_dir.mkdir(parents=True, exist_ok=True)
    print(f"Каталог вывода: {output_dir.resolve()}")

    journals = load_students(str(students_path))
    total = 0
    for jname, classes in journals:
        if not classes:
            continue
        print(f"Журнал: {jname}")
        make_journal(jname, classes, args.period, dates, output_dir)
        total += sum(len(names) for _, _, names in classes)
    print(f"Готово. Всего учащихся: {total}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
